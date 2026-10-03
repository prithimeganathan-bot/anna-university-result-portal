import csv
import io
import openpyxl
from django.contrib import messages
from django.shortcuts import render, redirect
from .models import Student, SubjectResult

def home(request):
    return render(request, 'portalapp/home.html')

def get_result(request):
    if request.method == 'POST':
        register_number = request.POST.get('register_number').strip().upper()
        dob = request.POST.get('dob')

        try:
            student = Student.objects.get(
                register_number=register_number,
                date_of_birth=dob
            )
            # Store student id in session
            request.session['student_id'] = student.id
            request.session['active_tab'] = 'profile'
            return redirect('dashboard')

        except Student.DoesNotExist:
            return render(request, 'portalapp/home.html', {
                'error': 'Invalid Register Number or Date of Birth.'
            })

    return redirect('home')


def dashboard(request):
    student_id = request.session.get('student_id')
    if not student_id:
        return redirect('home')

    student = Student.objects.get(id=student_id)
    tab = request.GET.get('tab', 'profile')

    subjects = student.subjects.all()
    total_subjects = subjects.count()
    passed = subjects.filter(result='Pass').count()
    failed = total_subjects - passed
    overall_result = 'PASS' if failed == 0 else 'FAIL'

    context = {
        'student': student,
        'subjects': subjects,
        'tab': tab,
        'overall_result': overall_result,
        'passed': passed,
        'failed': failed,
    }
    return render(request, 'portalapp/dashboard.html', context)


def logout_view(request):
    request.session.flush()
    return redirect('home')
def institution_login(request):
    if request.method == 'POST':
        code = request.POST.get('inst_code', '').strip()
        password = request.POST.get('inst_password', '').strip()

        # Any code/password works — it's a mimic
        if code and password:
            request.session['inst_code'] = code
            return redirect('inst_dashboard')
        else:
            return render(request, 'portalapp/home.html', {
                'inst_error': 'Invalid Institution Code or Password.'
            })
    return redirect('home')


def inst_dashboard(request):
    inst_code = request.session.get('inst_code', '3114')
    tab = request.GET.get('tab', 'home')
    context = {
        'inst_code': inst_code,
        'tab': tab,
    }
    return render(request, 'portalapp/inst_dashboard.html', context)


def inst_logout(request):
    request.session.flush()
    return redirect('home')

def bulk_upload(request):
    if request.method == 'POST' and request.FILES.get('student_file'):
        uploaded_file = request.FILES['student_file']
        filename = uploaded_file.name
        success_count = 0
        error_list = []

        try:
            # ── CSV handling ──
            if filename.endswith('.csv'):
                decoded = uploaded_file.read().decode('utf-8')
                reader = csv.DictReader(io.StringIO(decoded))
                rows = list(reader)

            # ── Excel handling ──
            elif filename.endswith('.xlsx') or filename.endswith('.xls'):
                wb = openpyxl.load_workbook(uploaded_file)
                ws = wb.active
                headers = [cell.value for cell in ws[1]]
                rows = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    rows.append(dict(zip(headers, row)))
            else:
                messages.error(request, 'Only .csv or .xlsx files are supported.')
                return render(request, 'portalapp/bulk_upload.html')

            for i, row in enumerate(rows, start=2):
                try:
                    # Skip empty rows
                    if not row.get('register_number'):
                        continue

                    # Create or update student
                    student, created = Student.objects.update_or_create(
                        register_number=str(row['register_number']).strip(),
                        defaults={
                            'name': str(row['name']).strip(),
                            'date_of_birth': str(row['date_of_birth']).strip(),
                            'department': str(row['department']).strip(),
                            'college_name': str(row['college_name']).strip(),
                            'semester': int(row['semester']),
                            'regulation': str(row['regulation']).strip(),
                        }
                    )

                    # Add subjects if provided
                    if row.get('subject_code'):
                        SubjectResult.objects.update_or_create(
                            student=student,
                            subject_code=str(row['subject_code']).strip(),
                            defaults={
                                'subject_name': str(row['subject_name']).strip(),
                                'internal_marks': int(row['internal_marks']),
                                'external_marks': int(row['external_marks']),
                                'total_marks': int(row['total_marks']),
                                'grade': str(row['grade']).strip(),
                                'result': str(row['result']).strip(),
                            }
                        )
                    success_count += 1

                except Exception as e:
                    error_list.append(f'Row {i}: {str(e)}')

            messages.success(request, f'✅ Successfully uploaded {success_count} records!')
            if error_list:
                for err in error_list:
                    messages.warning(request, err)

        except Exception as e:
            messages.error(request, f'File error: {str(e)}')

    return render(request, 'portalapp/bulk_upload.html')